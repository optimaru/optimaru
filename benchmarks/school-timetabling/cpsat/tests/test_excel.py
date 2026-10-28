"""Comcigan Excel import / assignment inspection / export."""

import csv
import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from teacher_scheduling import (
    ExcelFormatError,
    ExportOptions,
    NeisFormat,
    SolverConfig,
    import_workbook,
    inspect_timetable,
    neis_bundle,
    neis_zip,
    solve_partial,
    timetable_workbook,
    unassigned_courses,
)
from teacher_scheduling.domain.enums import Weekday
from teacher_scheduling.excel import SHEET_BASICS, SHEET_DAILY, SHEET_LOAD, SHEET_SPECIAL
from teacher_scheduling.sample_workbook import CURRICULUM, DAILY_PERIODS, write_workbook


@pytest.fixture(scope="module")
def sample_bytes() -> bytes:
    return write_workbook()


@pytest.fixture(scope="module")
def imported(sample_bytes: bytes):
    return import_workbook(io.BytesIO(sample_bytes))


@pytest.fixture(scope="module")
def solved(imported):
    return solve_partial(imported.problem, SolverConfig(time_limit_seconds=60))


# ---------------------------------------------------------------- import


def test_sample_workbook_imports_without_warnings(imported):
    assert imported.warnings == []
    assert imported.basics.school_name == "Comcigan Middle School"
    assert imported.basics.school_year == 2026
    assert imported.basics.classes_per_grade == {1: 3, 2: 3, 3: 3}


def test_calendar_is_the_longest_day_across_grades(imported):
    # Only grade 3 has 7 periods on Mon/Tue, so the school calendar has 7 periods too.
    assert imported.problem.calendar.periods_per_day[Weekday.MON] == 7
    assert imported.problem.calendar.periods_per_day[Weekday.FRI] == 6


def test_shorter_grades_block_the_extra_periods(imported):
    """Per-grade periods per weekday are expressed as class unavailability."""
    by_id = imported.problem.class_group_by_id
    first = by_id["C101"]
    third = by_id["C301"]
    blocked = {(slot.day, slot.period) for slot in first.unavailable}
    assert (Weekday.MON, 7) in blocked, "grade 1 has 6 periods on Monday"
    assert (Weekday.TUE, 7) in blocked
    assert third.unavailable == frozenset(), "grade 3 fills the calendar"


def test_weekly_hours_match_the_curriculum(imported):
    problem = imported.problem
    for group in problem.class_groups:
        assigned = sum(c.periods_per_week for c in problem.courses_by_class_group[group.id])
        assert assigned == sum(CURRICULUM[group.grade].values())
        assert assigned == sum(DAILY_PERIODS[group.grade].values())


def test_official_names_collapse_into_one_short_subject(imported):
    # TechHome = Technology and Home Economics. Scheduling uses the short name; the official one stays for NEIS.
    assert "TechHome" in imported.problem.subject_by_id
    assert imported.official_subject_names["TechHome"] == ["Technology and Home Economics"]


def test_special_sheet_applies_blocks_and_bans(imported):
    problem = imported.problem
    art = [c for c in problem.courses if c.subject_id == "Art"]
    assert art and all(c.block_periods == 2 for c in art), "consecutive lessons: 2-period blocks"

    banned = {t.name: {str(s) for s in t.unavailable} for t in problem.teachers if t.unavailable}
    assert banned == {"Kim Haneul": {"Mon1"}, "Kwon Nayun": {"Fri6"}}


def test_usage_hints_below_the_table_are_not_parsed(imported):
    """The usage notes below the special-operations table must not be read as data."""
    assert not any("How to use" in w for w in imported.warnings)


def test_subject_room_requirement_comes_from_subject_sheet(imported):
    assert imported.problem.subject_by_id["Science"].required_room_type == "science lab"
    assert imported.problem.subject_by_id["Korean"].required_room_type is None
    assert imported.problem.subject_by_id["PE"].avoided_periods == frozenset({1})
    assert imported.problem.subject_by_id["Math"].preferred_periods == frozenset({1, 2, 3, 4})


# ---------------------------------------------------------------- format errors


def _minimal_workbook(**sheets: list[list]) -> io.BytesIO:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def test_missing_required_sheet_is_reported():
    with pytest.raises(ExcelFormatError, match="required sheet"):
        import_workbook(_minimal_workbook(**{SHEET_BASICS: [["학교명", "A"]]}))


def test_missing_header_column_is_reported():
    source = _minimal_workbook(
        **{
            SHEET_BASICS: [["1학년 학급수", 1]],
            SHEET_DAILY: [["학년", "월"], ["1학년", 3]],
            SHEET_LOAD: [["순", "단축과목명", "교사명"]],
        }
    )
    with pytest.raises(ExcelFormatError, match="정식과목명"):
        import_workbook(source)


def test_capacity_mismatch_becomes_a_warning():
    """When the load total differs from the periods per weekday, warn up front like Comcigan."""
    source = _minimal_workbook(
        **{
            SHEET_BASICS: [["1학년 학급수", 1]],
            SHEET_DAILY: [["학년", "월", "화"], ["1학년", 3, 3]],
            SHEET_LOAD: [
                ["순", "정식과목명", "단축과목명", "교사명", "1학년", None],
                [None, None, None, None, 1, None],
                [1, "Korean", "Korean", "Teacher Kim", 4, None],
            ],
        }
    )
    result = import_workbook(source)
    assert any("2 cells stay empty" in w for w in result.warnings)


# ---------------------------------------------------------------- scheduling + inspection


def test_imported_problem_solves_completely(solved, imported):
    required = sum(c.periods_per_week for c in imported.problem.courses)
    assert solved.status.has_solution
    assert len(solved.lessons) == required
    assert unassigned_courses(imported.problem, solved) == []


def test_inspection_has_no_errors_on_a_full_solve(imported, solved):
    report = inspect_timetable(imported.problem, solved)
    assert report.hard_violations == ()
    errors = {item.id: item.count for item in report.items if item.severity == "error"}
    assert errors == {"unassigned": 0, "class_empty_period": 0}
    assert report.total_lessons == report.required_lessons


def test_inspection_reports_every_search_condition(imported, solved):
    report = inspect_timetable(imported.problem, solved)
    ids = {item.id for item in report.items}
    assert {
        "unassigned",
        "class_empty_period",
        "same_subject_twice_a_day",
        "teacher_daily_overload",
        "consecutive_long",
        "lunch_adjacent",
        "teacher_gap",
    } <= ids


def test_partial_solve_leaves_unassigned_instead_of_infeasible():
    """Unfillable periods stay unassigned and the rest is placed."""
    source = _minimal_workbook(
        **{
            SHEET_BASICS: [["1학년 학급수", 1]],
            SHEET_DAILY: [["학년", "월"], ["1학년", 2]],
            SHEET_LOAD: [
                ["순", "정식과목명", "단축과목명", "교사명", "1학년", None],
                [None, None, None, None, 1, None],
                # Monday has only 2 cells but 5 periods are required -> 3 unassigned
                [1, "Korean", "Korean", "Teacher Kim", 5, None],
            ],
        }
    )
    problem = import_workbook(source).problem
    timetable = solve_partial(problem, SolverConfig(time_limit_seconds=10))
    assert timetable.status.has_solution, "partial timetabling must not be INFEASIBLE"
    assert len(timetable.lessons) == 2
    missing = unassigned_courses(problem, timetable)
    assert [(u.subject_id, u.missing) for u in missing] == [("Korean", 3)]

    report = inspect_timetable(problem, timetable)
    assert next(item for item in report.items if item.id == "unassigned").count == 1


def _parallel_workbook(periods_per_week: int, periods_on_monday: int) -> io.BytesIO:
    """Minimal workbook grouping grade 1 class 1's technology / home economics into one parallel group."""
    return _minimal_workbook(
        **{
            SHEET_BASICS: [["1학년 학급수", 1]],
            SHEET_DAILY: [["학년", "월"], ["1학년", periods_on_monday]],
            SHEET_LOAD: [
                ["순", "정식과목명", "단축과목명", "교사명", "1학년", None],
                [None, None, None, None, 1, None],
                [1, "Tech", "Tech", "Kim", periods_per_week, None],
                [2, "HomeEc", "HomeEc", "Lee", periods_per_week, None],
            ],
            SHEET_SPECIAL: [
                ["구분", "대상", "값"],
                ["분반", "1학년 1반/Tech", "TechHomeA"],
                ["분반", "1학년 1반/HomeEc", "TechHomeA"],
            ],
        }
    )


def test_parallel_group_is_not_penalized_twice_for_one_missing_period():
    """Parallel sections always move together (H5), so the unassigned penalty counts once per group."""
    # Monday has only 2 cells, the two parallel courses need 4 each -> 2 periods unassigned per group
    problem = import_workbook(_parallel_workbook(periods_per_week=4, periods_on_monday=2)).problem
    timetable = solve_partial(problem, SolverConfig(time_limit_seconds=10))

    assert len(timetable.lessons) == 4, "the two parallel courses share the 2 cells"
    missing = {u.subject_id: u.missing for u in unassigned_courses(problem, timetable)}
    assert missing == {"Tech": 2, "HomeEc": 2}, "the unassigned list still shows each course"
    assert timetable.penalties["weekly_period_counts"] == 2, "must not be multiplied by the group size"


# ---------------------------------------------------------------- export


def test_neis_bundle_has_one_csv_per_class(imported, solved):
    files = neis_bundle(imported.problem, solved, ExportOptions(school_year=2026, semester=1))
    assert len(files) == len(imported.problem.class_groups)
    assert "기초시간표(2026-1,1-1).csv" in files

    rows = list(csv.reader(io.StringIO(files["기초시간표(2026-1,1-1).csv"].decode("cp949"))))
    assert rows[0] == ["교시", "월요일", "화요일", "수요일", "목요일", "금요일"]
    # Grade 1 has 6 periods on Monday -> the Monday cell of the period-7 row is empty.
    assert rows[7][1] == ""
    assert all("/" in cell for cell in rows[1][1:] if cell)


def test_neis_uses_official_subject_names(imported, solved):
    options = ExportOptions(
        school_year=2026,
        official_subject_names={k: v[0] for k, v in imported.official_subject_names.items()},
    )
    text = neis_bundle(imported.problem, solved, options)["기초시간표(2026-1,1-1).csv"].decode("cp949")
    assert "Technology and Home Economics" in text and "TechHome/" not in text


def test_neis_format_can_widen_to_seven_day_columns(imported, solved):
    """The NEIS screen has 7 columns (Mon–Sun), so the column count must be extendable."""
    options = ExportOptions(school_year=2026, neis=NeisFormat(day_count=7))
    text = neis_bundle(imported.problem, solved, options)["기초시간표(2026-1,1-1).csv"].decode("cp949")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0] == ["교시", "월요일", "화요일", "수요일", "목요일", "금요일", "토요일", "일요일"]
    # Sat/Sun without lessons are empty columns
    assert all(row[6] == "" and row[7] == "" for row in rows[1:])


def test_neis_format_can_split_the_teacher_into_its_own_column(imported, solved):
    options = ExportOptions(school_year=2026, neis=NeisFormat(teacher_column=True))
    text = neis_bundle(imported.problem, solved, options)["기초시간표(2026-1,1-1).csv"].decode("cp949")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0][:3] == ["교시", "월요일", "월요일 교사"]
    subject, teacher = rows[1][1], rows[1][2]
    assert subject and teacher and "/" not in subject


def test_neis_format_can_drop_the_header_and_period_suffix(imported, solved):
    options = ExportOptions(
        school_year=2026,
        neis=NeisFormat(include_header=False, period_suffix=""),
    )
    text = neis_bundle(imported.problem, solved, options)["기초시간표(2026-1,1-1).csv"].decode("cp949")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0][0] == "1", "without a header the first row is period 1"
    assert len(rows) == imported.problem.calendar.max_periods


@pytest.fixture(scope="module")
def parallel_solved():
    """Minimal timetable where two parallel courses share a cell."""
    problem = import_workbook(_parallel_workbook(periods_per_week=2, periods_on_monday=2)).problem
    return problem, solve_partial(problem, SolverConfig(time_limit_seconds=10))


def test_neis_writes_every_course_sharing_a_cell(parallel_solved):
    """Parallel sections in one cell must not silently shrink to one."""
    problem, timetable = parallel_solved
    files = neis_bundle(problem, timetable, ExportOptions(school_year=2026))
    cell = list(csv.reader(io.StringIO(files["기초시간표(2026-1,1-1).csv"].decode("cp949"))))[1][1]
    assert "Tech/Kim" in cell and "HomeEc/Lee" in cell
    assert cell.count(";") == 1


def test_neis_can_keep_only_one_course_per_cell(parallel_solved):
    """If NEIS reads only one subject per cell, keep a single representative."""
    problem, timetable = parallel_solved
    options = ExportOptions(school_year=2026, neis=NeisFormat(parallel_separator=None))
    text = neis_bundle(problem, timetable, options)["기초시간표(2026-1,1-1).csv"].decode("cp949")
    cell = list(csv.reader(io.StringIO(text)))[1][1]
    assert cell in {"Tech/Kim", "HomeEc/Lee"}


def test_neis_format_rejects_an_impossible_day_count():
    with pytest.raises(ValueError, match="day_count"):
        NeisFormat(day_count=8)


def test_neis_zip_contains_the_same_files(imported, solved):
    data = neis_zip(imported.problem, solved, ExportOptions(school_year=2026))
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert len(archive.namelist()) == len(imported.problem.class_groups)


def test_print_workbook_has_the_expected_sheets(imported, solved):
    data = timetable_workbook(imported.problem, solved, ExportOptions(school_year=2026))
    workbook = load_workbook(io.BytesIO(data))
    assert {"Overview", "Classes", "Teachers", "Special rooms", "Teacher load"} <= set(workbook.sheetnames)
    summary = workbook["Teacher load"]
    total = sum(summary.cell(row=r, column=3).value for r in range(2, summary.max_row + 1))
    assert total == sum(len(lesson.teacher_ids) for lesson in solved.lessons)
