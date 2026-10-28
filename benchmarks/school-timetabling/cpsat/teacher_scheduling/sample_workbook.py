"""컴시간 양식 샘플 엑셀 생성.

`excel.import_workbook` 이 읽는 양식 그대로의 예제 파일을 만든다. 사용자가
[교사별시수표 엑셀양식] 버튼으로 받아 자기 학교 자료로 덮어쓰는 용도이고,
그대로 올리면 풀리는 완결된 데이터이기도 하다.

학교 구성: 3개 학년 x 3개 반, 주 32~34시간, 특별실 5종.
"""

from __future__ import annotations

from io import BytesIO
from typing import IO, Any

GRADES = (1, 2, 3)
CLASSES_PER_GRADE = 3

#: 학년별 요일 시수. 3학년만 월/화가 한 시간 길다 — 학년별 시수 처리를 보여준다.
DAILY_PERIODS: dict[int, dict[str, int]] = {
    1: {"월": 6, "화": 6, "수": 7, "목": 7, "금": 6},
    2: {"월": 6, "화": 6, "수": 7, "목": 7, "금": 6},
    3: {"월": 7, "화": 7, "수": 7, "목": 7, "금": 6},
}

#: (단축과목명, 정식과목명, 특별실, 선호교시, 기피교시)
SUBJECTS: tuple[tuple[str, str, str, str, str], ...] = (
    ("국어", "국어", "", "", ""),
    ("수학", "수학", "", "1-4", ""),
    ("영어", "영어", "", "", ""),
    ("과학", "과학", "과학실", "", ""),
    ("사회", "사회", "", "", ""),
    ("체육", "체육", "체육관", "", "1"),
    ("음악", "음악", "음악실", "", ""),
    ("미술", "미술", "미술실", "", ""),
    ("도덕", "도덕", "", "", ""),
    ("기가", "기술·가정", "", "", ""),
    ("정보", "정보", "컴퓨터실", "", ""),
)

#: 학년별 과목 시수 (합계가 그 학년 주간 교시 수와 같아야 한다).
CURRICULUM: dict[int, dict[str, int]] = {
    1: {"국어": 5, "수학": 5, "영어": 5, "과학": 4, "사회": 3, "체육": 3, "음악": 2, "미술": 2, "도덕": 1, "기가": 2},
    2: {"국어": 5, "수학": 5, "영어": 5, "과학": 4, "사회": 3, "체육": 3, "음악": 2, "미술": 2, "도덕": 1, "기가": 2},
    3: {
        "국어": 5,
        "수학": 5,
        "영어": 5,
        "과학": 4,
        "사회": 4,
        "체육": 3,
        "음악": 2,
        "미술": 2,
        "도덕": 1,
        "기가": 2,
        "정보": 1,
    },
}

#: 과목 -> 교사 이름들. 담당 학급은 아래에서 균등 분배한다.
TEACHERS: dict[str, tuple[str, ...]] = {
    "국어": ("김하늘", "이서준", "박지현"),
    "수학": ("최민준", "정예린", "강도윤"),
    "영어": ("조은비", "윤태호", "장우진"),
    "과학": ("한소율", "서준호"),
    "사회": ("문가영", "배시우"),
    "체육": ("오지환", "임세리"),
    "음악": ("신하람",),
    "미술": ("황도경",),
    "도덕": ("권나윤",),
    "기가": ("남기훈",),
    "정보": ("탁현우",),
}

ROOMS: tuple[tuple[str, str], ...] = (
    ("과학실 1", "과학실"),
    ("과학실 2", "과학실"),
    ("체육관", "체육관"),
    ("음악실", "음악실"),
    ("미술실", "미술실"),
    ("컴퓨터실", "컴퓨터실"),
)

SPECIAL_OPERATIONS: tuple[tuple[str, str, str], ...] = (
    ("연속수업", "미술", "2"),
    ("배정금지", "김하늘", "월1"),
    ("배정금지", "권나윤", "금6"),
)


def _all_classes() -> list[tuple[int, int]]:
    return [(grade, number) for grade in GRADES for number in range(1, CLASSES_PER_GRADE + 1)]


def _assign_teacher(subject: str, grade: int, number: int) -> str:
    """과목 담당 교사들에게 학급을 돌아가며 나눠 준다."""
    names = TEACHERS[subject]
    index = (GRADES.index(grade) * CLASSES_PER_GRADE + (number - 1)) % len(names)
    return names[index]


def build_rows() -> list[tuple[str, str, str, dict[tuple[int, int], int]]]:
    """(정식과목명, 단축과목명, 교사명, {(학년,반): 시수}) 목록.

    컴시간 규칙대로 (정식과목, 단축과목, 교사) 가 같은 것만 한 줄로 묶는다.
    """
    official = {short: full for short, full, *_ in SUBJECTS}
    grouped: dict[tuple[str, str], dict[tuple[int, int], int]] = {}
    for grade, number in _all_classes():
        for subject, hours in CURRICULUM[grade].items():
            teacher = _assign_teacher(subject, grade, number)
            grouped.setdefault((subject, teacher), {})[grade, number] = hours
    return [
        (official[subject], subject, teacher, hours)
        for (subject, teacher), hours in sorted(grouped.items(), key=lambda item: (item[0][0], item[0][1]))
    ]


def write_workbook(destination: str | IO[bytes] | None = None) -> bytes:
    """샘플 워크북을 만들어 저장하고, 바이트도 함께 돌려준다."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="BBBBBB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor="EEF3FB")

    def style_header(cell: Any) -> None:
        cell.font = header_font
        cell.alignment = center
        cell.fill = fill
        cell.border = box

    # ---------------------------------------------------------- 기본사항
    basics = workbook.active
    basics.title = "기본사항"
    basics.column_dimensions["A"].width = 18
    basics.column_dimensions["B"].width = 24
    rows: list[tuple[str, Any]] = [
        ("항목", "값"),
        ("학교명", "컴시간중학교"),
        ("학년도", 2026),
        *[(f"{grade}학년 학급수", CLASSES_PER_GRADE) for grade in GRADES],
        ("점심시간 직전교시", 4),
    ]
    for index, (key, value) in enumerate(rows, start=1):
        basics.cell(row=index, column=1, value=key)
        basics.cell(row=index, column=2, value=value)
    for column in (1, 2):
        style_header(basics.cell(row=1, column=column))

    # ---------------------------------------------------------- 요일별시수
    daily = workbook.create_sheet("요일별시수")
    day_names = ["월", "화", "수", "목", "금"]
    daily.cell(row=1, column=1, value="학년")
    for index, day in enumerate(day_names, start=2):
        daily.cell(row=1, column=index, value=day)
    daily.cell(row=1, column=len(day_names) + 2, value="계")
    for offset, grade in enumerate(GRADES, start=2):
        daily.cell(row=offset, column=1, value=f"{grade}학년")
        for index, day in enumerate(day_names, start=2):
            daily.cell(row=offset, column=index, value=DAILY_PERIODS[grade][day])
        daily.cell(row=offset, column=len(day_names) + 2, value=sum(DAILY_PERIODS[grade].values()))
    for column in range(1, len(day_names) + 3):
        style_header(daily.cell(row=1, column=column))
        daily.column_dimensions[get_column_letter(column)].width = 8

    # ---------------------------------------------------------- 교사별시수표
    sheet = workbook.create_sheet("교사별시수표")
    fixed = ["순", "정식과목명", "단축과목명", "교사명"]
    sheet.cell(row=1, column=1, value="교사별 시수표").font = Font(bold=True, size=14)

    # 2행: 고정 열 + 학년 라벨(병합), 3행: 반 번호
    for index, label in enumerate(fixed, start=1):
        sheet.cell(row=2, column=index, value=label)
        sheet.merge_cells(start_row=2, start_column=index, end_row=3, end_column=index)
    column = len(fixed) + 1
    for grade in GRADES:
        start = column
        for number in range(1, CLASSES_PER_GRADE + 1):
            sheet.cell(row=3, column=column, value=number)
            column += 1
        sheet.cell(row=2, column=start, value=f"{grade}학년")
        sheet.merge_cells(start_row=2, start_column=start, end_row=2, end_column=column - 1)
    total_column = column
    sheet.cell(row=2, column=total_column, value="계")
    sheet.merge_cells(start_row=2, start_column=total_column, end_row=3, end_column=total_column)

    for row in (2, 3):
        for index in range(1, total_column + 1):
            style_header(sheet.cell(row=row, column=index))

    class_columns = {
        (grade, number): len(fixed) + 1 + GRADES.index(grade) * CLASSES_PER_GRADE + (number - 1)
        for grade, number in _all_classes()
    }
    for order, (official, short, teacher, hours) in enumerate(build_rows(), start=1):
        row = 3 + order
        sheet.cell(row=row, column=1, value=order)
        sheet.cell(row=row, column=2, value=official)
        sheet.cell(row=row, column=3, value=short)
        sheet.cell(row=row, column=4, value=teacher)
        for key, value in hours.items():
            sheet.cell(row=row, column=class_columns[key], value=value).alignment = center
        sheet.cell(row=row, column=total_column, value=sum(hours.values())).alignment = center
        for index in range(1, total_column + 1):
            sheet.cell(row=row, column=index).border = box
    for index in range(1, total_column + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 12 if index in (2, 3, 4) else 5

    # ---------------------------------------------------------- 과목
    subjects = workbook.create_sheet("과목")
    headers = ["단축과목명", "정식과목명", "특별실", "복수교사", "선호교시", "기피교시"]
    for index, label in enumerate(headers, start=1):
        style_header(subjects.cell(row=1, column=index, value=label))
        subjects.column_dimensions[get_column_letter(index)].width = 14
    for offset, (short, full, room, preferred, avoided) in enumerate(SUBJECTS, start=2):
        subjects.cell(row=offset, column=1, value=short)
        subjects.cell(row=offset, column=2, value=full)
        subjects.cell(row=offset, column=3, value=room)
        subjects.cell(row=offset, column=4, value="")
        subjects.cell(row=offset, column=5, value=preferred)
        subjects.cell(row=offset, column=6, value=avoided)

    # ---------------------------------------------------------- 특별실
    rooms_sheet = workbook.create_sheet("특별실")
    for index, label in enumerate(["특별실명", "종류"], start=1):
        style_header(rooms_sheet.cell(row=1, column=index, value=label))
        rooms_sheet.column_dimensions[get_column_letter(index)].width = 16
    for offset, (name, room_type) in enumerate(ROOMS, start=2):
        rooms_sheet.cell(row=offset, column=1, value=name)
        rooms_sheet.cell(row=offset, column=2, value=room_type)

    # ---------------------------------------------------------- 특별작업
    special = workbook.create_sheet("특별작업")
    for index, label in enumerate(["구분", "대상", "값"], start=1):
        style_header(special.cell(row=1, column=index, value=label))
        special.column_dimensions[get_column_letter(index)].width = 18
    for offset, (kind, target, value) in enumerate(SPECIAL_OPERATIONS, start=2):
        special.cell(row=offset, column=1, value=kind)
        special.cell(row=offset, column=2, value=target)
        special.cell(row=offset, column=3, value=value)
    special.cell(row=len(SPECIAL_OPERATIONS) + 3, column=1, value="구분 사용법:")
    for offset, hint in enumerate(
        (
            "배정금지 | 교사명 | 월1        - 그 교사가 그 시간에 수업 불가",
            "연속수업 | 단축과목명 | 2   - 그 과목을 2연강으로 붙임",
            "고정배정 | 1학년 1반/도덕 | 수3 - 그 강좌를 그 시간에 고정",
            "분반    | 1학년 1반/수학 | 수학A - 같은 값끼리 동시 진행",
        ),
        start=len(SPECIAL_OPERATIONS) + 4,
    ):
        special.cell(row=offset, column=1, value=hint)

    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    if destination is not None:
        if isinstance(destination, str):
            with open(destination, "wb") as handle:
                handle.write(data)
        else:
            destination.write(data)
    return data


__all__ = ["CURRICULUM", "DAILY_PERIODS", "build_rows", "write_workbook"]
